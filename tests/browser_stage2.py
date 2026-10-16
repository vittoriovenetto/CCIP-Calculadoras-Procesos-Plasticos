"""Etapa 2: operaciones reales y lectura independiente XLSX/PDF, sin publicar."""
import json
import os
import re
import unittest
from pathlib import Path
import openpyxl
import fitz
from pypdf import PdfReader
from browser_smoke import BrowserSmoke, VARIANTS, ARTIFACTS, ROOT

KEY='ccip-calculators-web-v1'
class Stage2(BrowserSmoke):
    def stored(self):
        return self.page.evaluate(f"JSON.parse(localStorage.getItem('{KEY}')||'[]')")

    def download(self, fmt, saved=False, index=0):
        with self.page.expect_download(timeout=90000) as event:
            if saved:
                self.page.locator('[data-history-export]').nth(index).select_option(fmt)
            else:
                self.page.locator('#export-format').select_option(fmt)
                self.page.locator('#export').click()
        target=ARTIFACTS/event.value.suggested_filename
        event.value.save_as(target)
        self.page.wait_for_function("!document.getElementById('export').disabled")
        return target

    def check_excel(self,path,snapshot,lang='es'):
        book=openpyxl.load_workbook(path,data_only=False)
        expected=['Resumen','Parámetros','Resultados','Fuentes y observaciones'] if lang=='es' else ['Summary','Parameters','Results','Sources and observations']
        self.assertEqual(book.sheetnames[:4],expected)
        for sheet in book:
            self.assertEqual(sheet.freeze_panes,'A2')
            for row in sheet:
                for cell in row:self.assertNotEqual(cell.data_type,'f','Un texto se convirtió en fórmula')
        params=list(book.worksheets[1].values)[1:]
        for row,(key,value) in zip(params,snapshot['inputs'].items()):
            if value not in ['',None] and isinstance(value,(float,int)) or isinstance(value,str) and re.fullmatch(r'[+-]?\d+(\.\d+)?',value):
                self.assertIsInstance(row[1],(int,float));self.assertAlmostEqual(row[1],float(value),places=10)
            if value is not None and value!='':self.assertEqual(str(row[4]),str(value))
        for row,metric in zip(list(book.worksheets[2].values)[1:],snapshot.get('metrics',[])):
            if metric['value'] is None:self.assertIn(row[1],['No disponible','Unavailable'])
            else:self.assertAlmostEqual(row[1],metric['value'],places=9)
            unit=metric.get('unit')
            if lang=='en':unit={'moneda/h':'currency/h','moneda/pza':'currency/part','moneda/kg':'currency/kg','pzas/h':'parts/h','ciclos/h':'cycles/h','envases/h':'containers/h'}.get(unit,unit)
            self.assertEqual(row[2],unit or ('No disponible' if lang=='es' else 'Unavailable'))
        return book

    def check_pdf(self,path,language='es'):
        reader=PdfReader(path)
        self.assertGreater(len(reader.pages),0)
        joined='\n'.join(p.extract_text() for p in reader.pages).replace('\u200b','')
        self.assertIn('CCIP',joined)
        self.assertNotIn('\ufffd',joined)
        doc=fitz.open(path)
        for i,page in enumerate(doc):
            self.assertAlmostEqual(page.rect.width,595.28,delta=.1);self.assertAlmostEqual(page.rect.height,841.89,delta=.1)
            self.assertIn(f"{'Página' if language=='es' else 'Page'} {i+1} / {len(doc)}",page.get_text())
            for word in page.get_text('words'):
                self.assertGreaterEqual(word[0],34);self.assertLessEqual(word[2],page.rect.width-34)
                self.assertGreaterEqual(word[1],30);self.assertLessEqual(word[3],page.rect.height-20)
        return joined,len(reader.pages)

    def test_S01_new_overwrite_restore_cancel_delete_ids(self):
        self.page.locator('[data-field="partMass"]').fill('31.125')
        self.save_case('Duplicado');first=self.stored()[0]
        self.save_case('Duplicado');second=self.stored()[0]
        self.assertNotEqual(first['id'],second['id'])
        self.assertNotEqual(self.page.locator('#compare-a').input_value(),self.page.locator('#compare-b').input_value(),'La selección automática debe comparar registros distintos')
        self.page.locator('[data-history]').nth(1).click()
        self.page.locator('[data-field="partMass"]').fill('55.125')
        before=self.stored();self.page.locator('#calculate').click();self.assertEqual(self.stored(),before)
        self.page.locator('#save').click();self.page.locator('#save-mode').select_option('overwrite')
        self.assertEqual(self.page.locator('#save-target').input_value(),first['id'])
        self.assertTrue(self.page.locator('#save-confirm').is_disabled())
        self.page.locator('#save-overwrite-confirm').check();self.page.locator('#save-cancel').click();self.assertEqual(self.stored(),before)
        self.page.locator('#save').click();self.page.locator('#save-mode').select_option('overwrite');self.page.locator('#save-overwrite-confirm').check();self.page.locator('#save-confirm').click()
        after=self.stored();self.assertEqual(len(after),2);self.assertEqual(after[0],second)
        self.assertEqual(after[1]['id'],first['id']);self.assertEqual(after[1]['created_at'],first['created_at']);self.assertEqual(after[1]['inputs']['partMass'],'55.125');self.assertNotEqual(after[1]['updated_at'],first['updated_at'])
        self.save_case('Copia');self.assertEqual(len(self.stored()),3)
        self.page.locator('#compare-a').select_option(first['id']);self.page.locator('#compare-b').select_option(second['id'])
        self.page.once('dialog',lambda d:d.dismiss());self.page.locator(f'[data-history-delete="{first["id"]}"]').click();self.assertEqual(len(self.stored()),3)
        self.page.once('dialog',lambda d:(self.assertIn('Duplicado',d.message),d.accept()))
        self.page.locator(f'[data-history-delete="{first["id"]}"]').click()
        self.assertEqual(self.page.locator('#compare-a').input_value(),'');self.assertEqual(self.page.locator('#compare-b').input_value(),second['id'])
        self.assertEqual(self.stored()[1],second)
        self.select_process('costos');self.page.locator('#save').click()
        self.assertTrue(self.page.locator('#save-mode option[value="overwrite"]').evaluate('el=>el.disabled'));self.page.keyboard.press('Escape')

    def test_S02_legacy_migration_corruption_and_storage_failures(self):
        legacy={'module':'inyeccion','level':'technical','date':'2026-01-02T03:04:05Z','inputs':{'partMass':'012.34567890','unknown':'kept'},'metrics':[{'label':'Histórico','value':987.654321,'unit':'g','formula':'m'}],'warnings':['No recalcular'],'future':{'keep':True}}
        self.page.evaluate('(data)=>localStorage.setItem("'+KEY+'",JSON.stringify(data))',[legacy,legacy,None]);self.page.reload()
        data=self.stored();self.assertNotEqual(data[0]['id'],data[1]['id']);self.assertEqual(data[0]['metrics'],legacy['metrics']);self.assertEqual(data[0]['inputs'],legacy['inputs']);self.assertIsNone(data[2])
        self.page.reload();self.assertEqual(self.stored(),data)
        self.page.evaluate("()=>{Storage.prototype.setItem=function(){throw new DOMException('Quota','QuotaExceededError')}}")
        self.page.on('dialog',lambda d:d.accept())
        self.page.locator('[data-history-delete]').nth(0).click();self.assertEqual(self.stored(),data);self.assertEqual(self.page.locator('[data-history]').count(),2)
        self.save_case('Fallará');self.assertTrue(self.page.locator('#save-dialog').is_visible());self.assertEqual(self.stored(),data);self.page.keyboard.press('Escape')
        self.page.reload();self.page.evaluate('(key)=>localStorage.setItem(key,"{")',KEY);self.page.reload()
        self.page.locator('#save').click();self.assertEqual(self.page.evaluate('(key)=>localStorage.getItem(key)',KEY),'{');self.page.locator('#calculate').click();self.assertGreater(self.page.locator('.kpi').count(),0)

    def test_S03_active_and_saved_real_exports_all_processes_languages(self):
        requested=[];self.page.on('request',lambda req:requested.append(req.url))
        self.page.reload();self.assertFalse(any('/vendor/' in url for url in requested))
        for language in ['es','en']:
            self.page.locator('#language').select_option(language)
            for module,values in VARIANTS[:7]:
                self.fill_case(module,values,'advanced')
                if module=='inyeccion':
                    self.page.locator('#profile-details > summary').click();self.page.locator('#calculate').click()
                self.save_case(f'{module}-{language}')
                record=self.stored()[0]
                excel=self.download('xlsx');self.check_excel(excel,record,language)
                pdf=self.download('pdf');text,_=self.check_pdf(pdf,language)
                self.assertIn('cm³' if module=='inyeccion' else 'CCIP',text)
                saved_xlsx=self.download('xlsx',saved=True);self.check_excel(saved_xlsx,record,language)
                saved_pdf=self.download('pdf',saved=True);self.check_pdf(saved_pdf,language)
                js=json.loads(self.download('json',saved=True).read_text());self.assertEqual(js['inputs'],record['inputs']);self.assertEqual(js['metrics'],record['metrics']);self.assertEqual(js['export_metadata']['status'],'saved')
                if module=='inyeccion' and language=='es' and os.environ.get('CCIP_WRITE_EXAMPLES')=='1':
                    samples=ROOT/'docs/ejemplos/etapa-2';samples.mkdir(parents=True,exist_ok=True)
                    (samples/'inyeccion-ejemplo-academico.xlsx').write_bytes(saved_xlsx.read_bytes());(samples/'inyeccion-ejemplo-academico.pdf').write_bytes(saved_pdf.read_bytes())
        self.assertTrue(all(url.startswith(self.url.rsplit('/',2)[0]) or '/CCIP-' in url for url in requested),'Petición externa')

    def test_S04_snapshot_precision_injection_text_and_long_unicode_report(self):
        warning='Advertencia técnica: estimación geométrica, no validación industrial. '
        snapshot={'module':'inyeccion','level':'advanced','date':'2020-01-02T03:04:05Z','name':'=HYPERLINK("https://example.org","España")','inputs':{'partMass':'0001.234567890123','runnerMass':'','cavities':'0'},'metrics':[{'label':'Resultado histórico inmutable','value':987.6543210123,'unit':'cm³','formula':'π × D² / 4; Δsᵢ = 1000 × Vᵢ / A; ρₘ ≤ ρₛ; µm; °C; −'}],'material':{'source':{'id':'sintetica','url':'https://example.org/ficha','note':'EJEMPLO SINTÉTICO, sin datos industriales'}},'warnings':[warning*160,'FIN DE ADVERTENCIAS'],'notes':['No se inventan datos ausentes.']}
        self.page.evaluate('(data)=>localStorage.setItem("'+KEY+'",JSON.stringify([data]))',snapshot);self.page.reload();saved=self.stored()[0]
        self.page.locator('[data-field="partMass"]').fill('9999')
        excel=self.download('xlsx',True);book=self.check_excel(excel,saved)
        self.assertEqual(book['Resumen']['B2'].value,snapshot['name']);self.assertEqual(book['Resumen']['B2'].data_type,'s')
        self.assertIn('https://example.org/ficha',str(list(book.worksheets[3].values)))
        pdf=self.download('pdf',True);text,pages=self.check_pdf(pdf)
        self.assertGreater(pages,1);self.assertIn('FIN DE ADVERTENCIAS',text);self.assertIn('Resultado histórico inmutable',' '.join(text.split()));self.assertIn('ρₘ',text);self.assertIn('Δsᵢ',text);self.assertIn('https://example.org/ficha',text)
        js=json.loads(self.download('json',True).read_text());self.assertEqual(js['metrics'],snapshot['metrics']);self.assertEqual(js['inputs'],snapshot['inputs']);self.assertEqual(self.stored()[0],saved)
        active=json.loads(self.download('json').read_text());self.assertEqual(active['inputs']['partMass'],'9999');self.assertEqual(active['export_metadata']['status'],'active');self.assertEqual(self.stored()[0],saved)

    def test_S06_family_density_sources_profile_and_history_limit(self):
        self.density('HDPE','advanced')
        self.page.locator('#profile-details > summary').click()
        self.save_case('Perfil HDPE — ejemplo, no validado')
        saved=self.stored()[0]
        self.assertEqual(saved['density_assessment']['origin'],'family_reference')
        path=self.download('xlsx',True);book=self.check_excel(path,saved)
        self.assertIn('Perfil L1 L2 L3',book.sheetnames)
        self.assertEqual(book.worksheets[-1]['B2'].value,30)
        text=str(list(book.worksheets[3].values))
        self.assertIn('reiloyusa.com',text);self.assertIn('reference_generic_estimate',text)
        pdf=self.download('pdf',True);text,_=self.check_pdf(pdf)
        self.assertIn('Estimación con densidad de familia',text)
        if os.environ.get('CCIP_WRITE_EXAMPLES')=='1':
            samples=ROOT/'docs/ejemplos/etapa-2';samples.mkdir(parents=True,exist_ok=True)
            (samples/'inyeccion-ejemplo-academico.xlsx').write_bytes(path.read_bytes())
            (samples/'inyeccion-ejemplo-academico.pdf').write_bytes(pdf.read_bytes())
        # Un nombre igual a una etiqueta traducible sigue siendo un nombre propio.
        self.page.once('dialog',lambda d:d.accept('Pendiente'))
        self.page.locator('[data-history-rename]').nth(0).click()
        self.page.locator('#language').select_option('en')
        self.assertIn('Pendiente',self.page.locator('#compare-a option').nth(1).inner_text())
        self.assertIn('Pendiente',self.page.locator('#history').inner_text())
        records=[dict(saved,id=f'case-{i}',name=f'Case {i}') for i in range(60)]
        self.page.evaluate('(records)=>localStorage.setItem("'+KEY+'",JSON.stringify(records))',records)
        self.page.reload();self.assertEqual(self.page.locator('[data-history]').count(),15)
        for count in [30,45,60]:
            self.page.locator('#history-more').click();self.assertEqual(self.page.locator('[data-history]').count(),count)
        before=self.stored();messages=[]
        self.page.on('dialog',lambda d:(messages.append(d.message),d.accept()))
        self.save_case('61');self.assertEqual(self.stored(),before);self.assertIn('60',messages[-1]);self.page.keyboard.press('Escape')

    def test_S05_responsive_dialog_history_and_export_failures(self):
        self.save_case('Historial de pruebas — datos sintéticos')
        for language in ['es','en']:
            self.page.locator('#language').select_option(language)
            for width in [360,390,540,768,1024,1440]:
                self.page.set_viewport_size({'width':width,'height':1000})
                self.assert_layout();self.page.locator('#save').click();self.assert_layout()
                self.page.screenshot(path=str(ARTIFACTS/f'etapa2-guardar-{language}-{width}.png'))
                self.page.keyboard.press('Escape');self.page.locator('#history').scroll_into_view_if_needed();self.assert_layout()
                self.page.screenshot(path=str(ARTIFACTS/f'etapa2-historial-{language}-{width}.png'))
        self.page.route('**/vendor/exceljs*',lambda route:route.abort())
        errors=[];self.page.on('dialog',lambda d:(errors.append(d.message),d.accept()))
        self.page.locator('#export-format').select_option('xlsx');self.page.locator('#export').click()
        self.page.wait_for_function("!document.getElementById('export').disabled")
        self.assertEqual(errors,['The file could not be generated.']);self.assertEqual(len(self.stored()),1)
        # El fallo deliberado de red es esperado, sin errores JS.
        self.errors=[e for e in self.errors if 'ERR_FAILED' not in e]

if __name__=='__main__':
    suite=unittest.TestSuite(Stage2(name) for name in unittest.defaultTestLoader.getTestCaseNames(Stage2) if name.startswith('test_S'))
    result=unittest.TextTestRunner(verbosity=2).run(suite)
    raise SystemExit(not result.wasSuccessful())
