"""Stage 4 browser checks. GitHub receipt API is explicitly mocked, no public test Issue."""
import json
import unittest
import openpyxl
from pypdf import PdfReader
from browser_stage2 import Stage2, ARTIFACTS
from browser_smoke import MATERIALS

class Stage4(Stage2):
    def explore(self,key=''):
        self.page.goto(self.url+'?view=materials'+('&material='+key if key else ''))
        self.page.locator('#explorer-count').filter(has_text='resultados').wait_for()
    def new_material(self):
        self.page.locator('#explorer-new').click()
        form=self.page.locator('#custom-form')
        form.locator('[name=family]').select_option('PP')
        form.locator('[name=grade]').fill('Synthetic laboratory fixture, not a real grade')
        form.locator('[name=source]').fill('Synthetic browser source')
        form.locator('[name=provenance]').fill('Controlled browser fixture, no manufacturer claim')
        form.locator('[name=material_kind]').select_option('recycled')
        form.locator('[name=verification_declaration]').check()
        form.locator('button[type=submit]').click()
        self.page.locator('#explorer-edit').wait_for()
        return self.page.evaluate("JSON.parse(localStorage.getItem('ccip-user-materials-v1'))[0]")
    def filter(self,key,value):
        self.page.locator('#explorer-controls details').evaluate('(el)=>el.open=true')
        self.page.locator(f'[data-filter={key}]').select_option(value)
    def test_F01_navigation_history_preserves_results_and_draft(self):
        self.page.locator('#field-partMass').fill('43.125')
        self.page.locator('#calculate').click()
        result=self.page.locator('#summary-metrics').inner_text()
        self.page.locator('#custom-panel').evaluate('(el)=>el.open=true')
        self.page.locator('#custom-form [name=grade]').fill('Unsaved draft')
        self.page.locator('#materials-page').click();self.page.locator('#explorer-query').wait_for()
        self.page.locator('#explorer-query').fill('terluran GP-35')
        self.assertIn('1 resultados',self.page.locator('#explorer-count').inner_text())
        self.page.locator('[data-entry="catalog-abs-terluran-gp-35"]').click()
        self.assertIn('material=catalog-abs-terluran-gp-35',self.page.url)
        self.page.locator('#learn-page').click();self.page.locator('#learning-back').wait_for()
        self.page.go_back();self.page.locator('#explorer-ficha h2').wait_for()
        self.page.locator('#materials-back').click()
        self.assertEqual(self.page.locator('#field-partMass').input_value(),'43.125')
        self.assertEqual(self.page.locator('#summary-metrics').inner_text(),result)
        self.assertEqual(self.page.locator('#custom-form [name=grade]').input_value(),'Unsaved draft')
        self.page.go_back();self.page.locator('#explorer-ficha h2').wait_for()
        self.page.locator('#language').select_option('en');self.page.locator('#materials-view h1').filter(has_text='Materials library').wait_for()
    def test_F02_all_original_fichas_multiple_flows_experiments_and_filters(self):
        self.explore()
        self.assertIn('88 resultados',self.page.locator('#explorer-count').inner_text())
        self.filter('family','HDPE');self.filter('manufacturer','Dow');self.filter('category','commercial');self.filter('property','mfr')
        self.assertGreater(self.page.locator('.explorer-card').count(),0)
        for card in self.page.locator('.explorer-card').all():self.assertIn('HDPE · Dow',card.inner_text())
        self.page.locator('#explorer-clear').click()
        ids=[m['id'] for m in MATERIALS]
        for key in ids:
            self.page.locator('#materials-page').click()
            self.page.locator('#explorer-query').wait_for()
            self.page.evaluate('(key)=>{const a=document.createElement("a");a.dataset.entry=key;a.href="?view=materials&material="+key;document.querySelector("#explorer-list").append(a);a.click()}',key)
            self.page.locator('#explorer-ficha h2').wait_for()
            self.assertIn(key,self.page.locator('#explorer-ficha').inner_text())
            original=next(m for m in MATERIALS if m['id']==key)
            self.assertEqual(self.page.locator('#explorer-ficha .pvt-chart').count(),int(original.get('pvt',{}).get('status')=='grade_exact_chart_available_approx_digitization'))
        self.explore('catalog-abs-terluran-gp-35')
        text=self.page.locator('#explorer-ficha').inner_text();self.assertIn('MFR',text);self.assertIn('MVR',text);self.assertIn('ASTM D1238',text);self.assertIn('ISO 1133',text)
        self.explore();self.filter('category','experimental')
        self.assertIn('15 resultados',self.page.locator('#explorer-count').inner_text())
        self.page.locator('[data-entry]').first.click();self.assertEqual(self.page.locator('#explorer-use').count(),0);self.assertIn('experiment',self.page.locator('#explorer-ficha').inner_text().lower())
    def test_F03_comparison_and_real_exports(self):
        self.explore();self.page.locator('#explorer-query').fill('DOW')
        self.page.locator('[data-compare]').nth(0).check();self.page.locator('[data-compare]').nth(1).check()
        self.assertGreater(self.page.locator('#explorer-comparison [data-status]').count(),0)
        for fmt in ['xlsx','pdf','json']:
            with self.page.expect_download(timeout=90000) as event:self.page.locator(f'[data-export=comparison][data-format={fmt}]').click()
            path=ARTIFACTS/event.value.suggested_filename;event.value.save_as(path)
            if fmt=='xlsx':
                book=openpyxl.load_workbook(path);self.assertEqual(book.sheetnames[0],'Identidad y estado');self.assertGreater(book.worksheets[1].max_row,2)
                for sheet in book:
                    for row in sheet:
                        for cell in row:self.assertNotEqual(cell.data_type,'f')
            elif fmt=='pdf':self.assertIn('CCIP',' '.join(p.extract_text() for p in PdfReader(path).pages))
            else:self.assertEqual(len(json.loads(path.read_text())['entries']),2)
        self.page.locator('#explorer-clear-comparison').click();self.assertEqual(self.page.locator('[data-compare]:checked').count(),0)
        self.page.locator('[data-entry]').first.click()
        with self.page.expect_download(timeout=90000) as event:self.page.locator('[data-export=ficha][data-format=xlsx]').click()
        path=ARTIFACTS/'ficha.xlsx';event.value.save_as(path);self.assertEqual(openpyxl.load_workbook(path).sheetnames[0],'Identidad y estado')
    def test_F04_local_crud_receipt_mock_and_no_implicit_submission(self):
        self.explore();m=self.new_material();self.assertEqual(m['revision'],1)
        self.assertTrue(all(v is None for v in m['properties'].values()))
        self.page.locator('#explorer-edit').click();form=self.page.locator('#custom-form');form.locator('[name=lot]').fill('Synthetic lot B');form.locator('button[type=submit]').click()
        edit=self.page.evaluate("JSON.parse(localStorage.getItem('ccip-user-materials-v1'))[0]")
        self.assertEqual(edit['id'],m['id']);self.assertEqual(edit['revision'],2)
        self.page.locator('#contribution-panel').evaluate('(el)=>el.open=true')
        self.page.locator('#contribution-reason').fill('Controlled fixture, not a real public submission')
        self.page.locator('#contribution-prepare').click();self.page.locator('#contribution-open').wait_for()
        prepared=self.page.evaluate("JSON.parse(localStorage.getItem('ccip-material-proposals-v1'))[0]")
        self.assertFalse(prepared['remote_confirmed']);self.assertEqual(prepared['state'],'prepared')
        body=self.page.locator('#contribution-body').input_value();url='https://github.com/vittoriovenetto/CCIP-Calculadoras-Procesos-Plasticos/issues/99999'
        self.page.route('https://api.github.com/repos/vittoriovenetto/CCIP-Calculadoras-Procesos-Plasticos/issues/99999',lambda route:route.fulfill(json={'number':99999,'html_url':url,'body':body,'labels':[],'state':'open'}))
        self.page.locator('#contribution-url').fill(url);self.page.locator('#contribution-confirm').click();self.page.wait_for_function("JSON.parse(localStorage.getItem('ccip-material-proposals-v1'))[0].remote_confirmed")
        self.assertIn('pendiente de revisión',self.page.locator('#contribution-status').inner_text().lower())
        self.page.reload();self.page.locator('#explorer-edit').wait_for();self.assertEqual(self.page.evaluate("JSON.parse(localStorage.getItem('ccip-user-materials-v1')).length"),1)
        self.page.locator('#explorer-use').click();self.assertEqual(self.page.locator('#material-grade').input_value(),m['id']);self.assertEqual(self.page.locator('#field-meltDensity').input_value(),'')
        self.page.locator('#custom-panel').evaluate('(el)=>el.open=true');self.page.once('dialog',lambda d:d.accept());self.page.locator('[data-personal-delete]').click();self.assertEqual(self.page.evaluate("JSON.parse(localStorage.getItem('ccip-user-materials-v1')).length"),0)
    def test_F06_shared_fixture_conflict_snapshots_and_update_proposal(self):
        fixture=self.page.evaluate('''async()=>{
            const {makeCustomMaterial}=await import('./src/custom-materials.js');
            const {emptySharedCatalog}=await import('./src/shared-materials.js');
            const approved=makeCustomMaterial({family:'PP',grade:'Approved synthetic fixture, not a real grade',source:'Synthetic test source',provenance:'Controlled browser test, not an actual approved contribution',verification_declaration:true,material_kind:'virgin',melt_density:.81,melt_density_conditions:'220 °C; 5 MPa'});
            const local=structuredClone(approved);local.grade='Edited local synthetic fixture';local.revision=2;
            localStorage.setItem('ccip-user-materials-v1',JSON.stringify([local]));
            const url='https://github.com/vittoriovenetto/CCIP-Calculadoras-Procesos-Plasticos/issues/99999',pr='https://github.com/vittoriovenetto/CCIP-Calculadoras-Procesos-Plasticos/pull/99999',date='2026-10-16T12:00:00.000Z';
            const entry={id:approved.id,material:approved,category:'Agregado por usuarios',revision:1,incorporated_at:date,updated_at:date,verification_status:'approved_contribution_not_industrially_certified',contribution:{url},review:{state:'approved_as_contribution',url:pr,reviewer:'synthetic-fixture',changed_fields:['material'],justification:'Test fixture only'},history:[{revision:1,date,previous:null,proposed:approved,changed_fields:['material'],justification:'Test fixture only',source:url,review:pr}]};
            return {approved,local,shared:{...emptySharedCatalog(),revision:2,materials:[entry]}};
        }''')
        self.page.route('**/user-materials.v1.json',lambda route:route.fulfill(json=fixture['shared']))
        self.explore(fixture['approved']['id']);self.page.locator('#explorer-use').click()
        self.assertEqual(self.page.locator('#field-meltDensity').input_value(),'0.81')
        self.save_case('Shared synthetic test snapshot');saved=self.stored()[0]
        self.assertEqual(saved['material']['grade_snapshot']['catalog_origin'],'shared')
        self.assertEqual(saved['material']['grade_snapshot']['shared_revision'],1)
        self.explore(fixture['local']['id']+'~local');self.assertEqual(self.page.locator('#explorer-use').count(),0)
        self.assertIn('Conflicto',self.page.locator('#explorer-ficha').inner_text())
        self.page.locator('#explorer-edit').click();self.page.locator('#custom-form [name=lot]').fill('Changed synthetic lot')
        self.page.locator('#custom-form button[type=submit]').click()
        self.assertIn('~local',__import__('urllib.parse',fromlist=['unquote']).unquote(self.page.url))
        self.page.locator('#contribution-panel').evaluate('(el)=>el.open=true')
        self.page.locator('#contribution-base').click();self.page.locator('#contribution-reason').fill('Synthetic update fixture')
        self.page.locator('#contribution-prepare').click();self.page.locator('#contribution-open').wait_for()
        payload=self.page.evaluate("JSON.parse(localStorage.getItem('ccip-material-proposals-v1'))[0].payload")
        self.assertEqual(payload['action'],'update');self.assertEqual(payload['base_revision'],1);self.assertEqual(payload['material']['id'],fixture['approved']['id'])
        self.assertEqual(self.stored()[0],saved)
        self.page.locator('#explorer-copy').click()
        locals=self.page.evaluate("JSON.parse(localStorage.getItem('ccip-user-materials-v1'))")
        self.assertEqual(len(locals),2);self.assertEqual(len(set(x['id'] for x in locals)),2)
        self.assertIn('copia local',self.page.locator('#explorer-ficha h2').inner_text())
    def test_F07_html_is_text_and_prepared_proposal_can_resume_after_reload(self):
        self.explore();m=self.new_material();self.page.locator('#explorer-edit').click()
        literal='<img src=x onerror="globalThis.ccipXSS=1">'
        self.page.locator('#custom-form [name=grade]').fill(literal);self.page.locator('#custom-form button[type=submit]').click()
        self.assertEqual(self.page.locator('#explorer-ficha h2').inner_text(),literal)
        self.assertIsNone(self.page.evaluate('globalThis.ccipXSS'))
        self.page.locator('#contribution-panel').evaluate('(el)=>el.open=true');self.page.locator('#contribution-reason').fill('Explicit synthetic public-content preview, never actually sent')
        self.page.locator('#contribution-prepare').click();self.page.locator('#contribution-open').wait_for()
        first=self.page.evaluate("JSON.parse(localStorage.getItem('ccip-material-proposals-v1'))[0]")
        self.page.reload();self.page.locator('#contribution-panel').evaluate('(el)=>el.open=true')
        self.page.locator('#contribution-earlier').select_option(first['payload']['submission_id']);self.page.locator('#contribution-resume').click();self.page.locator('#contribution-open').wait_for()
        self.assertIn(first['payload']['submission_id'],self.page.locator('#contribution-body').input_value())
        self.assertFalse(self.page.evaluate("JSON.parse(localStorage.getItem('ccip-material-proposals-v1'))[0].remote_confirmed"))
        self.assertIsNone(self.page.evaluate('globalThis.ccipXSS'))
    def test_F05_responsive_bilingual_screenshots_no_page_overflow(self):
        self.explore('catalog-abs-terluran-gp-35')
        for lang in ['es','en']:
            self.page.locator('#language').select_option(lang)
            for width in [360,390,540,768,1024,1440]:
                self.page.set_viewport_size({'width':width,'height':900})
                self.page.wait_for_timeout(50)
                self.assertTrue(self.page.evaluate('document.documentElement.scrollWidth <= innerWidth + 1'),f'{lang} {width}')
                self.page.screenshot(path=str(ARTIFACTS/f'etapa-4-{lang}-{width}.png'),full_page=True)
                self.page.locator('#materials-page').click();self.page.locator('#explorer-query').wait_for();self.page.screenshot(path=str(ARTIFACTS/f'biblioteca-{lang}-{width}.png'))
                self.page.locator('#explorer-query').fill('Terluran GP-35');self.page.locator('[data-entry="catalog-abs-terluran-gp-35"]').click()
                self.assertEqual(self.page.locator('#explorer-ficha h2').inner_text(),'Terluran GP-35')
        self.page.locator('#materials-back').click();self.assertTrue(self.page.locator('#calculator-view').is_visible())

if __name__=='__main__':
    suite=unittest.TestSuite(Stage4(name) for name in unittest.defaultTestLoader.getTestCaseNames(Stage4) if name.startswith('test_F'))
    result=unittest.TextTestRunner(verbosity=2).run(suite)
    raise SystemExit(not result.wasSuccessful())
