"""UX-01–UX-20 and routing/draft regressions. Local app only; no publishing."""
import json
import os
import unittest
from pathlib import Path
import openpyxl
from browser_stage6 import Stage6, ARTIFACTS

WIDTHS=(320,360,390,540,768,1024,1440,1920)

class MachineExplorerUX(Stage6):
    def library(self, machine=None):
        self.page.locator('#machines-page').click()
        self.page.locator('[data-machine-card]').first.wait_for()
        if machine:
            self.page.locator(f'[data-machine-card="{machine}"] [data-machine-route]').first.click()
            self.page.locator(f'[data-machine-detail="{machine}"]').wait_for()

    def machine_filter(self,key,value):
        self.page.locator('#machine-filters').evaluate('el=>el.open=true')
        self.page.locator(f'[data-machine-filter={key}]').select_option(value)

    def compare_machine(self,id):
        self.page.locator(f'[data-machine-card="{id}"] [data-machine-compare]').click()

    def report_download(self,target,format):
        with self.page.expect_download(timeout=90000) as event:
            self.page.locator(f'[data-machine-export={target}][data-format={format}]').click()
        path=ARTIFACTS/event.value.suggested_filename
        event.value.save_as(path)
        return path

    def test_UX01_material_action_below_grade_and_three_columns_stack_on_mobile(self):
        for width in (390,1440):
            self.page.set_viewport_size({'width':width,'height':1000})
            grade=self.page.locator('#material-grade').bounding_box()
            action=self.page.locator('#custom-material-toggle').bounding_box()
            self.assertAlmostEqual(grade['x'],action['x'],delta=2)
            self.assertGreaterEqual(action['y'],grade['y']+grade['height'])
            if width==390:
                selectors=[self.page.locator('#material-'+name).bounding_box() for name in ('family','manufacturer','grade')]
                self.assertTrue(selectors[0]['y']<selectors[1]['y']<selectors[2]['y'])
            self.assert_layout()

    def test_UX02_editor_full_width_keyboard_draft_survives_navigation(self):
        self.page.locator('#custom-material-toggle').focus();self.page.keyboard.press('Enter')
        form=self.page.locator('#custom-form')
        form.locator('[name=grade]').fill('UX unsaved synthetic draft')
        self.page.evaluate("document.dispatchEvent(new Event('ccip-language-change'))")
        self.assertEqual(self.page.evaluate('document.activeElement.name'),'grade')
        for width in WIDTHS:
            self.page.set_viewport_size({'width':width,'height':1000})
            self.assertGreater(self.page.locator('#custom-material-editor').bounding_box()['width'],self.page.locator('#materials').bounding_box()['width']-40)
            self.assert_layout()
        self.library();self.page.locator('#machines-back').click()
        self.assertEqual(form.locator('[name=grade]').input_value(),'UX unsaved synthetic draft')
        self.assertEqual(self.page.locator('#custom-material-toggle').get_attribute('aria-expanded'),'true')

    def test_UX03_create_edit_import_counter_and_select_local_material(self):
        self.page.locator('#custom-material-toggle').click()
        form=self.page.locator('#custom-form');form.locator('[name=family]').select_option('PP')
        form.locator('[name=grade]').fill('UX local fixture');form.locator('button[type=submit]').click()
        self.assertIn('(1)',self.page.locator('#custom-material-toggle').inner_text())
        id=self.page.locator('#material-grade').input_value();self.assertTrue(id.startswith('user-'))
        self.page.locator('#custom-material-toggle').click();self.page.locator('[data-personal-edit]').click()
        form.locator('[name=grade]').fill('UX edited fixture');form.locator('button[type=submit]').click()
        self.assertEqual(self.page.locator('#material-grade').input_value(),id)
        self.page.locator('#materials-page').click();self.page.locator('#explorer-new').click()
        form.locator('[name=grade]').fill('UX library fixture');form.locator('button[type=submit]').click()
        self.page.locator('#explorer-use').wait_for();self.page.locator('#materials-back').click()
        self.assertIn('(2)',self.page.locator('#custom-material-toggle').inner_text())
        added=self.page.evaluate("JSON.parse(localStorage.getItem('ccip-user-materials-v1')).find(m=>m.grade==='UX library fixture')")
        self.page.locator('#material-family').select_option(added['family'])
        self.page.locator('#material-manufacturer').select_option(added['manufacturer'])
        self.page.locator('#material-grade').select_option(added['id'])
        self.assertEqual(self.page.locator('#material-grade').input_value(),added['id'])
        self.assert_layout()

    def test_UX04_optional_equipment_is_compact_without_selection(self):
        self.page.locator('#machine-select option').nth(1).wait_for(state='attached')
        for width in WIDTHS:
            self.page.set_viewport_size({'width':width,'height':1000})
            panel=self.page.locator('#machine-panel')
            self.assertLess(panel.bounding_box()['height'],250 if width<=540 else 165)
            self.assertEqual(panel.locator('details[open]').count(),0)
            self.assertEqual(self.page.locator('#machine-select').input_value(),'')
            self.assert_layout()

    def test_UX05_selected_equipment_summary_and_datasheet_actions(self):
        self.equipment('moiron-m2-i')
        self.assertTrue(self.page.locator('.machine-application-status').is_visible())
        self.assertFalse(self.page.locator('#machine-datasheet').evaluate('el=>el.open'))
        self.assertIn('Moiron',self.page.locator('#machine-panel').inner_text())
        self.page.locator('.machine-reference-actions a').click()
        self.page.locator('[data-machine-detail="moiron-m2-i"]').wait_for()
        self.page.locator('#machines-back').click()
        self.assertEqual(self.page.locator('#machine-select').input_value(),'moiron-m2-i')

    def test_UX06_automatic_values_without_confirmation_checkbox(self):
        self.equipment('arburg-allrounder-370s-700-290-s30')
        self.assertEqual(self.page.locator('#machine-confirm').count(),0)
        self.assertEqual(self.page.locator('#field-screwD').input_value(),'30')
        self.assertEqual(self.page.locator('#field-screwMaxStroke').input_value(),'150')
        for width in (390,1440):
            self.page.set_viewport_size({'width':width,'height':1000});self.page.locator('#machine-panel').evaluate('el=>el.scrollIntoView({block:"start"})')
            self.page.screenshot(path=str(ARTIFACTS/f'ux-automatic-machine-es-{width}.png'))
        self.page.locator('#language').select_option('en')
        self.assertIn('Selected machine specification',self.page.locator('#provenance-screwD').inner_text())

    def test_UX07_manual_machine_inputs_preserved(self):
        self.page.locator('#field-screwD').fill('17')
        self.equipment('moiron-m2-i')
        self.assertEqual(self.page.locator('#field-screwD').input_value(),'17')
        self.assertEqual(self.page.locator('#field-screwMaxStroke').input_value(),'')

    def test_UX08_switch_removes_only_old_automatic_values(self):
        self.equipment('arburg-allrounder-370s-700-290-s30')
        self.page.locator('#machine-select').select_option('moiron-m2-i')
        self.assertEqual(self.page.locator('#field-screwD').input_value(),'20')
        self.assertEqual(self.page.locator('#field-screwMaxStroke').input_value(),'')

    def test_UX09_remove_selection_keeps_manual_values(self):
        self.equipment('arburg-allrounder-370s-700-290-s30')
        self.page.locator('#field-screwD').fill('35');self.page.locator('#field-partMass').fill('13.125')
        self.page.locator('.machine-reference-actions button').last.click()
        self.assertEqual(self.page.locator('#machine-select').input_value(),'')
        self.assertEqual(self.page.locator('#field-screwD').input_value(),'35')
        self.assertEqual(self.page.locator('#field-screwMaxStroke').input_value(),'')
        self.assertEqual(self.page.locator('#field-partMass').input_value(),'13.125')

    def test_UX10_no_calculator_menu_or_hidden_dependency(self):
        self.assertEqual(self.page.locator('#calculator-page').count(),0)
        self.assertEqual(self.page.locator('.learning-nav > a').count(),3)
        self.assertEqual(self.page.locator('.learning-nav b').all_text_contents(),['Materiales','Maquinaria','Aprendizaje'])
        self.assertNotIn('Calculadora',self.page.locator('.learning-nav').inner_text())

    def test_UX11_default_view_is_calculator(self):
        self.assertTrue(self.page.locator('#calculator-view').is_visible())
        self.assertTrue(self.page.locator('#machines-view').is_hidden())
        self.assertTrue(self.page.locator('#materials-view').is_hidden())
        self.assertTrue(self.page.locator('#learning-view').is_hidden())

    def test_UX12_three_sections_and_process_navigation(self):
        self.page.locator('#materials-page').click();self.page.locator('#explorer-query').wait_for()
        self.library();self.page.locator('#learn-page').click();self.page.locator('#learning-back').wait_for()
        self.select_process('pelicula')
        self.assertTrue(self.page.locator('#calculator-view').is_visible())
        self.assertEqual(self.page.locator('#process-select').input_value(),'pelicula')
        self.assertIn('process=pelicula',self.page.url)

    def test_UX13_logo_contextual_return_back_forward_preserve_case(self):
        self.page.locator('#field-partMass').fill('47.125');self.page.locator('#calculate').click()
        metrics=self.page.locator('#summary-metrics').inner_text()
        self.library('moiron-m2-i');self.page.locator('.brand').click()
        self.assertEqual(self.page.locator('#field-partMass').input_value(),'47.125')
        self.assertEqual(self.page.locator('#summary-metrics').inner_text(),metrics)
        self.page.go_back();self.page.locator('[data-machine-detail="moiron-m2-i"]').wait_for()
        self.page.go_forward();self.assertTrue(self.page.locator('#calculator-view').is_visible())

    def test_UX14_search_by_model_maker_process_partial_and_case(self):
        self.library();query=self.page.locator('#machine-query')
        query.fill('MOIRON');self.assertEqual(self.page.locator('[data-machine-card]').count(),3)
        query.fill('XH-430B-25')
        self.assertEqual(set(self.page.locator('[data-machine-card]').evaluate_all('els=>els.map(el=>el.dataset.machineCard)')),{'xihua-xh-430b-25','xihua-xh-430b-25-3','dongguan-xihua-xh430b-2-2kw'})
        query.fill('película XH430B');self.assertGreaterEqual(self.page.locator('[data-machine-card]').count(),2)
        query.fill('no-such-machine');self.assertEqual(self.page.locator('[data-machine-card]').count(),0)
        self.page.locator('#machine-clear-filters').click();self.assertEqual(self.page.locator('[data-machine-card]').count(),25)

    def test_UX15_all_six_filters_and_mobile_panel(self):
        self.page.set_viewport_size({'width':390,'height':1000});self.library()
        self.assertFalse(self.page.locator('#machine-filters').evaluate('el=>el.open'))
        for key,value in [('process','pelicula'),('manufacturer','Xihua'),('type','laboratory_blown_film'),('scope','specific_configuration'),('documentation','document_checked'),('specification','die_diameter_mm')]:self.machine_filter(key,value)
        self.assertEqual(self.page.locator('[data-machine-card]').count(),2)
        self.page.locator('#machine-clear-filters').click();self.machine_filter('scope','series_range')
        self.assertEqual(self.page.locator('[data-machine-card]').count(),2)
        self.assert_layout()

    def test_UX16_priority_datasheets_exact_variants_sources_and_missing_data(self):
        self.library('moiron-m2-i')
        self.assertIn('20 mm',self.page.locator('#machine-ficha').inner_text())
        self.assertIn('microplast.com',self.page.locator('#machine-ficha .machine-sources a').first.get_attribute('href'))
        self.page.locator('.machine-missing').evaluate('el=>el.open=true')
        self.assertIn('No documentado',self.page.locator('.machine-missing').inner_text())
        self.page.goto(self.url+'?view=machines&machine=xihua-xh-430b-25')
        self.page.locator('[data-machine-detail="xihua-xh-430b-25"]').wait_for()
        self.assertIn('25 mm',self.page.locator('#machine-ficha').inner_text())
        self.assertIn('/cmjxl/6.html',self.page.locator('#machine-ficha .machine-sources a').first.get_attribute('href'))

    def test_UX17_compare_two_to_four_and_export_excel(self):
        self.library()
        ids=['arburg-allrounder-370s-700-290-s30','moiron-m2-i','moiron-m2-ii','xihua-xh436-25t']
        for id in ids:self.compare_machine(id)
        self.assertEqual(self.page.locator('#machine-cards [data-machine-compare]:checked').count(),4)
        self.assertEqual(self.page.locator('.machine-pair').count(),6)
        self.compare_machine('moiron-page104-16mm-100kn')
        self.assertFalse(self.page.locator('[data-machine-card="moiron-page104-16mm-100kn"] input').is_checked())
        self.assertIn('cuatro',self.page.locator('#machine-library-status').inner_text())
        path=self.report_download('comparison','xlsx');book=openpyxl.load_workbook(path)
        self.assertGreaterEqual(book.worksheets[0].max_row,20)
        for sheet in book:
            for row in sheet:
                for cell in row:self.assertNotEqual(cell.data_type,'f')
        self.page.locator('#machine-clear-comparison').click();self.assertEqual(self.page.locator('[data-machine-compare]:checked').count(),0)

    def test_UX18_use_opens_correct_process_preserves_manual_inputs_and_returns_to_ficha(self):
        self.page.locator('#field-screwD').fill('17');self.page.locator('#language').select_option('en')
        self.library('xihua-xh-430b-25');self.page.locator('#machine-ficha [data-machine-use]').click()
        self.assertEqual(self.page.locator('#process-select').input_value(),'pelicula')
        self.assertEqual(self.page.locator('#field-dieDiameter').input_value(),'25')
        self.assertEqual(self.page.locator('#language').input_value(),'en')
        self.page.go_back();self.page.locator('[data-machine-detail="xihua-xh-430b-25"]').wait_for()
        self.page.locator('[data-machine-card="moiron-m2-i"] [data-machine-route]').first.click()
        self.page.locator('#machine-ficha [data-machine-use]').click()
        self.assertEqual(self.page.locator('#process-select').input_value(),'inyeccion')
        self.assertEqual(self.page.locator('#field-screwD').input_value(),'17')
        self.assertEqual(self.page.locator('#machine-select').input_value(),'moiron-m2-i')

    def test_UX19_images_without_permission_are_not_requested_and_fallback_is_functional(self):
        remote=[];self.page.on('request',lambda request:remote.append(request.url) if request.resource_type=='image' and not request.url.startswith(self.url) else None)
        self.library();self.assertEqual(self.page.locator('#machine-cards img').count(),0)
        self.assertEqual(self.page.locator('#machine-cards .machine-image-placeholder').count(),25)
        self.assertEqual(remote,[])

    def test_UX20_history_identity_snapshot_and_three_report_formats(self):
        self.equipment('moiron-m2-i');self.page.locator('#field-partMass').fill('13.125');self.save_case('UX immutable machine fixture')
        record=self.stored()[0]
        self.page.locator('[data-history-machine]').first.click();self.page.locator('[data-machine-detail="moiron-m2-i"]').wait_for()
        data=json.loads(self.report_download('datasheet','json').read_text())
        self.assertEqual(data['schema'],'ccip-machine-report/1');self.assertEqual(data['machines'][0]['id'],'moiron-m2-i')
        self.assertIn('microplast.com',str(data['machines'][0]['sources']))
        pdf=self.report_download('datasheet','pdf');self.assertTrue(pdf.read_bytes().startswith(b'%PDF'))
        text,_=self.check_pdf(pdf);self.assertIn('Moiron',text)
        book=openpyxl.load_workbook(self.report_download('datasheet','xlsx'));self.assertGreater(book.worksheets[1].max_row,4)
        self.assertEqual(self.stored()[0],record)
        self.page.locator('#machines-back').click();self.page.locator('[data-history]').first.click()
        self.assertEqual(self.stored()[0],record)

    def test_UX21_direct_urls_reload_and_invalid_machine_fallback(self):
        self.page.goto(self.url+'?view=machines&machine=xihua-xh-430b-25')
        self.page.locator('[data-machine-detail="xihua-xh-430b-25"]').wait_for();self.page.reload()
        self.page.locator('[data-machine-detail="xihua-xh-430b-25"]').wait_for()
        self.page.locator('#machine-ficha [data-machine-use]').click();self.page.locator('#level').select_option('advanced')
        self.page.reload();self.page.locator('#field-dieDiameter').wait_for()
        self.assertEqual(self.page.locator('#field-dieDiameter').input_value(),'25')
        self.assertEqual(self.page.locator('#level').input_value(),'advanced')
        self.page.goto(self.url+'?view=machines&machine=unknown-model')
        self.page.locator('#machine-ficha .warning').wait_for();self.page.locator('#machines-back').click()
        self.assertTrue(self.page.locator('#calculator-view').is_visible())

    def test_UX22_eight_widths_bilingual_real_screenshots_and_no_overflow(self):
        for lang in ('es','en'):
            self.page.locator('#language').select_option(lang)
            for width in WIDTHS:
                self.page.set_viewport_size({'width':width,'height':1000})
                self.assert_layout()
                self.library();self.assert_layout()
                if width in (390,1440):
                    self.page.locator('#machines-view').evaluate('el=>el.scrollIntoView({block:"start"})')
                    self.page.screenshot(path=str(ARTIFACTS/f'ux-machinery-library-{lang}-{width}.png'))
                self.page.locator('[data-machine-card="moiron-m2-i"] [data-machine-route]').first.click();self.assert_layout()
                if width in (390,1440):
                    self.page.locator('.machine-detail-heading').evaluate('el=>el.scrollIntoView({block:"start"})')
                    self.page.screenshot(path=str(ARTIFACTS/f'ux-moiron-ficha-{lang}-{width}.png'))
                self.page.locator('[data-machine-detail="moiron-m2-i"] [data-machine-compare]').click()
                self.compare_machine('moiron-m2-ii');self.assert_layout()
                if width in (390,1440):
                    self.page.locator('#machine-library-comparison').evaluate('el=>el.scrollIntoView({block:"start"})')
                    self.page.screenshot(path=str(ARTIFACTS/f'ux-machine-comparison-{lang}-{width}.png'))
                self.page.locator('#machine-clear-comparison').click()
                self.page.locator('[data-machine-card="xihua-xh-430b-25"] [data-machine-route]').first.click()
                if width in (390,1440):
                    self.page.locator('.machine-detail-heading').evaluate('el=>el.scrollIntoView({block:"start"})')
                    self.page.screenshot(path=str(ARTIFACTS/f'ux-xh-ficha-{lang}-{width}.png'))
                self.page.locator('#machines-back').click();self.assert_layout()
                if width in (390,1440):
                    self.page.locator('aside').evaluate('el=>el.scrollIntoView({block:"start"})')
                    self.page.screenshot(path=str(ARTIFACTS/f'ux-calculator-compact-{lang}-{width}.png'))
                    self.page.locator('#custom-material-toggle').click();self.assert_layout()
                    self.page.locator('#custom-form').evaluate('el=>el.scrollIntoView({block:"start"})')
                    self.page.screenshot(path=str(ARTIFACTS/f'ux-material-editor-{lang}-{width}.png'))
                    self.page.locator('#custom-material-toggle').click()

    def test_UX23_approved_test_image_failure_falls_back_without_layout_shift(self):
        record={'machine_id':'moiron-m2-i','manufacturer':'Moiron','source_type':'manufacturer_product_page','verified_at':'2026-10-17T12:00:00Z','permission':{'status':'approved','basis':'Synthetic test fixture only, not a product image licence'},'holder':'Test fixture','alt_es':'Prueba técnica','alt_en':'Technical test','source_url':'https://example.com/test','local_path':'./assets/machines/synthetic-test.webp','sha256':'a'*64}
        self.page.route('**/machine-images.v1.json',lambda route:route.fulfill(json={'schema':'ccip-machine-images/1','images':{'moiron-m2-i':record}}))
        self.page.route('**/assets/machines/synthetic-test.webp',lambda route:route.fulfill(status=200,content_type='image/webp',body=b'invalid test image'))
        with self.page.expect_request(lambda request:request.url.endswith('/assets/machines/synthetic-test.webp')):
            self.library()
            self.page.locator('[data-machine-card="moiron-m2-i"] .machine-image').scroll_into_view_if_needed()
        self.page.wait_for_function("document.querySelector('[data-machine-card=\"moiron-m2-i\"] .machine-image-placeholder') !== null && document.querySelector('[data-machine-card=\"moiron-m2-i\"] img') === null")
        self.page.wait_for_load_state('networkidle')
        box=self.page.locator('[data-machine-card="moiron-m2-i"] .machine-image').bounding_box()
        self.assertAlmostEqual(box['width']/box['height'],1.6,delta=.02)
        self.assert_layout()

    def test_UX24_unavailable_catalog_keeps_manual_calculation(self):
        self.page.route('**/machines.v1.json',lambda route:route.fulfill(status=200,content_type='application/json',body='{bad'))
        self.page.reload();self.page.locator('#field-partMass').fill('19');self.page.locator('#calculate').click()
        self.assertIn('19',self.page.locator('#summary-metrics').inner_text())
        self.page.locator('#machines-page').click();self.assertIn('no disponible',self.page.locator('#machine-library-status').inner_text())
        self.page.locator('#machines-back').click();self.assertEqual(self.page.locator('#field-partMass').input_value(),'19')

    def test_UX25_reload_preserves_per_tab_fields_editor_and_manual_origin_without_recalculating_history(self):
        self.equipment('arburg-allrounder-370s-700-290-s30');self.page.locator('#level').select_option('advanced')
        self.page.locator('#field-screwD').fill('35');self.page.locator('#field-partMass').fill('47.125')
        self.save_case('UX saved baseline');record=self.stored()[0]
        self.page.locator('#custom-material-toggle').click();self.page.locator('#custom-form [name=grade]').fill('UX draft after reload')
        self.page.reload();self.page.locator('#field-screwD').wait_for()
        self.assertEqual(self.page.locator('#field-screwD').input_value(),'35')
        self.assertEqual(self.page.locator('#field-partMass').input_value(),'47.125')
        self.assertEqual(self.page.locator('#machine-select').input_value(),'arburg-allrounder-370s-700-290-s30')
        self.assertEqual(self.page.locator('#custom-form [name=grade]').input_value(),'UX draft after reload')
        self.assertTrue(self.page.locator('#custom-panel').evaluate('el=>el.open'))
        self.assertEqual(self.stored()[0],record)
        self.page.locator('#calculate').click();self.save_case('UX new after reload')
        self.assertEqual(self.stored()[0]['calculation_metadata']['input_origins']['screwD'],'manual')
        self.assertEqual(self.stored()[1],record)


if __name__=='__main__':
    names=[name for name in unittest.defaultTestLoader.getTestCaseNames(MachineExplorerUX) if name.startswith('test_UX')]
    result=unittest.TextTestRunner(verbosity=2).run(unittest.TestSuite(MachineExplorerUX(name) for name in names))
    if os.environ.get('GITHUB_ACTIONS')=='true':
        for case,trace in result.failures+result.errors:
            message=trace[-12000:].replace('%','%25').replace('\r','%0D').replace('\n','%0A')
            print(f'::error file=tests/browser_machine_explorer.py,title={case.id()}::{message}',flush=True)
    raise SystemExit(not result.wasSuccessful())
